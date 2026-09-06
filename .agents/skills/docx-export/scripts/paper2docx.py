#!/usr/bin/env python3
"""把 Typst / LaTeX 论文源码导出为华为杯正文格式的 Word DOCX。

流程：
1. 识别 paper/ 下的 main.typ 或 main.tex，按 #include / \\input 顺序收集正文章节。
2. 把正文里引用的 PDF 图转成同名 PNG（Word 不支持 PDF 图片，PNG 写入 dpi 以保持物理尺寸）。
3. 调用 pandoc（typst/latex reader -> docx writer），公式输出为可编辑 OMML。
4. 用 python-docx 按 `_references/huaweibei_body_format.md` 套用官方正文格式：
   A4 四边 2.5 cm、宋体/Times New Roman 12 pt 首行缩进 2 字符、一级标题“一、”居中黑体 14 pt、
   二/三级 “1.1”/“1.1.1” 12 pt、图题下/表题上 11 pt 加粗、三线表、公式居中编号右对齐、
   摘要页、目录页（TOC 域 + 静态缓存）、页脚居中 9 pt 页码。封面不在本脚本范围（年年变，由人处理）。
5. 若有 soffice，顺手转一份 PDF 用于渲染抽检、统计页数，并回填目录页码。

用法：
    python paper2docx.py --paper paper --out paper/main.docx --pdf
    python paper2docx.py --paper paper --title "论文题目" --abstract-file abstract.txt \\
        --keywords "关键词1;关键词2;关键词3" --cover-placeholder
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

try:
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips
except ImportError:  # pragma: no cover
    sys.exit("缺少 python-docx：pip install python-docx")

# --------------------------------------------------------------------------- 常量（官方模板正文格式）
PAGE_W, PAGE_H = 11906, 16838          # A4，twips
MARGIN = 1418                          # 2.5 cm
HEADER_DIST, FOOTER_DIST = 851, 992
TEXT_W = PAGE_W - 2 * MARGIN           # 9070 版心宽
SONG, HEI, LI = "宋体", "黑体", "隶书"
TNR = "Times New Roman"
CODE_FONT = "Courier New"

CN_DIGITS = "零一二三四五六七八九"

PDF2PNG_LUA = r"""
-- 把 image 的 .pdf 源改为同名 .png（由 paper2docx.py 预先生成）
function Image(el)
  if el.src:match("%.pdf$") or el.src:match("%.PDF$") then
    el.src = el.src:gsub("%.[pP][dD][fF]$", ".png")
  end
  return el
end

local PAGE_BREAK = pandoc.RawBlock("openxml", '<w:p><w:r><w:br w:type="page"/></w:r></w:p>')

-- typst #pagebreak() -> Div.page-break；latex \newpage / \clearpage -> RawBlock；本脚本生成的摘要页 -> Div.pagebreak
function Div(el)
  if el.classes:includes("page-break") or el.classes:includes("pagebreak") then
    return PAGE_BREAK
  end
end

function RawBlock(el)
  if el.format == "latex" and (el.text:match("\\newpage") or el.text:match("\\clearpage")) then
    return PAGE_BREAK
  end
end

-- 图/表题自动编号：官方格式“图N xxx” / “表N xxx”（编号紧跟“图/表”，其后一个空格）
local fig_n, tab_n = 0, 0

local function prefix_caption(cap, label)
  if cap == nil or cap.long == nil or #cap.long == 0 then
    return cap
  end
  local first = cap.long[1]
  if first.t == "Plain" or first.t == "Para" then
    first.content:insert(1, pandoc.Str(label))
    first.content:insert(2, pandoc.Space())
  end
  return cap
end

function Figure(el)
  fig_n = fig_n + 1
  el.caption = prefix_caption(el.caption, "图" .. fig_n)
  return el
end

function Table(el)
  tab_n = tab_n + 1
  el.caption = prefix_caption(el.caption, "表" .. tab_n)
  return el
end

-- 参考文献 / 附录 一级标题不编号（Typst 的 numbering: none 不会传给 pandoc）
function Header(el)
  local txt = pandoc.utils.stringify(el)
  if el.level == 1 and (txt:match("^参考文献") or txt:match("^附录")) then
    el.classes:insert("unnumbered")
  end
  return el
end

-- 行间公式全篇连续编号；LaTeX \eqref{eq:x} / \ref{eq:x} 解析为 (N)
local eq_labels, eq_n = {}, 0

local function collect_math(el)
  if el.mathtype == "DisplayMath" then
    eq_n = eq_n + 1
    for lab in el.text:gmatch("\\label{([^}]+)}") do
      eq_labels[lab] = eq_n
    end
  end
  return nil
end

local function rewrite_ref(el)
  local ref = el.attributes["reference"]
  if ref and eq_labels[ref] then
    return pandoc.Str("(" .. eq_labels[ref] .. ")")
  end
end

-- 只含空 Span 的段落（Typst <label> 残留）删除
local function drop_empty(el)
  if #el.content == 0 then
    return {}
  end
  for _, x in ipairs(el.content) do
    if not (x.t == "Span" and #x.content == 0) then
      return nil
    end
  end
  return {}
end

return {
  { Math = collect_math },
  { Image = Image, Div = Div, RawBlock = RawBlock, Figure = Figure, Table = Table,
    Header = Header, Link = rewrite_ref, Para = drop_empty, Plain = drop_empty },
}
"""

TITLE_MARK = "\u200b\u200bTITLE\u200b\u200b"
ABSTRACT_MARK = "\u200b\u200bABSTRACT\u200b\u200b"
KEYWORDS_MARK = "\u200b\u200bKEYWORDS\u200b\u200b"

INCLUDE_TYP = re.compile(r'#include\s*\(?\s*"([^"]+)"\s*\)?')
INCLUDE_TEX = re.compile(r'\\(?:input|include)\{([^}]+)\}')
IMAGE_TYP = re.compile(r'image\(\s*"([^"]+\.pdf)"', re.I)
IMAGE_TEX = re.compile(r'\\includegraphics(?:\[[^\]]*\])?\{([^}]+\.pdf)\}', re.I)
PLACEHOLDER = re.compile(r'\[(论文标题|学校名称|参赛队号|成员 ?[A-C]|关键词\d?|中文摘要内容[^\]]*|封面[^\]]*)\]')
INTERNAL_NAMES = ("reports/", "figures/", "RESULTS_REPORT", "ANALYSIS_MODELING_REPORT", "CLAUDE.md", "AGENTS.md", "plan.md", "todo.md")
LABEL_TYP = re.compile(r'<([A-Za-z][\w:.-]*)>')
DISPLAY_EQ_TYP = re.compile(r'\$\s(?:\\\$|[^$])*?\s\$')
LEAD_PHRASE = re.compile(r'^(针对问题[一二三四五六七八九十\d]+[，,：:])')


@dataclass
class FrontMatter:
    title: str = ""
    abstract: list[str] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)


def run(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, text=True, capture_output=True, **kw)


def cn_number(n: int) -> str:
    """1 -> 一, 10 -> 十, 12 -> 十二, 21 -> 二十一（一级标题编号用）。"""
    if n <= 0 or n >= 100:
        return str(n)
    if n < 10:
        return CN_DIGITS[n]
    tens, ones = divmod(n, 10)
    return ("" if tens == 1 else CN_DIGITS[tens]) + "十" + (CN_DIGITS[ones] if ones else "")


# --------------------------------------------------------------------------- 源码收集
def detect_engine(paper: Path) -> tuple[str, Path]:
    if (paper / "main.typ").exists():
        return "typst", paper / "main.typ"
    if (paper / "main.tex").exists():
        return "latex", paper / "main.tex"
    sys.exit(f"{paper} 下找不到 main.typ 或 main.tex")


def collect_sections(engine: str, main: Path) -> list[Path]:
    text = main.read_text(encoding="utf-8")
    pattern = INCLUDE_TYP if engine == "typst" else INCLUDE_TEX
    if engine == "latex" and "\\begin{document}" in text:
        # 序言区里的 \input 通常是宏定义的附录 / 参考文献，排到正文之后
        preamble, body = text.split("\\begin{document}", 1)
        text = body + "\n" + preamble
    out: list[Path] = []
    for m in pattern.finditer(text):
        rel = m.group(1)
        p = main.parent / rel
        if engine == "latex" and not p.suffix:
            p = p.with_suffix(".tex")
        if p.exists():
            out.append(p)
        else:
            print(f"[warn] include 目标不存在，跳过: {rel}")
    return out


def extract_front_matter(engine: str, main: Path) -> FrontMatter:
    """从模板 main 文件中尽力抽取标题 / 摘要 / 关键词；抽不到就留空让 CLI 参数补。"""
    fm = FrontMatter()
    text = main.read_text(encoding="utf-8")
    if engine == "typst":
        m = re.search(r'#set document\(\s*title:\s*"([^"]*)"', text)
        if m:
            fm.title = m.group(1)
        lines = text.splitlines()
        # 摘要若被拆到单独文件（include("abstract.typ")），先内联进来再抽取
        for m_inc in re.finditer(r'include\("([^"]+)"\)', text):
            inc = main.parent / m_inc.group(1)
            if inc.exists() and '关键词：' in inc.read_text(encoding='utf-8'):
                idx = next(i for i, l in enumerate(lines) if m_inc.group(0) in l)
                lines = lines[:idx] + inc.read_text(encoding='utf-8').splitlines() + lines[idx + 1:]
                break
        code_lines = [(i, l) for i, l in enumerate(lines) if not l.lstrip().startswith('//')]
        start = next((i for i, l in code_lines if '要：' in l and '摘' in l), None)
        end = next((i for i, l in code_lines if '关键词：' in l), None)
        if start is not None and end is not None and end > start:
            noise = re.compile(r'^\s*(v\(|#?block\[|#set\b|#import\b|#v\(|pagebreak|\]\s*$)')
            kept = [l for l in lines[start + 1:end] if not noise.match(l)]
            body = '\n'.join(l.strip() for l in kept).strip()
            if body.startswith('[') and body.endswith(']'):
                body = body[1:-1].strip()
            fm.abstract = [p.strip() for p in re.split(r'\n\s*\n', body) if p.strip()]
            kw_line = lines[end].split('关键词：]', 1)[-1]
            fm.keywords = [k.strip() for k in re.findall(r'\[([^\]]+)\]', kw_line) if k.strip()]
            if not fm.keywords:
                fm.keywords = [k.strip() for k in re.split(r'#h\([^)]*\)|[;；]', kw_line) if k.strip()]
    else:
        m = re.search(r'\\newcommand\{\\papertitle\}\{(.+?)\}\s*$', text, re.M) \
            or re.search(r'\\heiti\\bfseries\s*(.+?)\}%', text)
        if m:
            fm.title = m.group(1).strip()
        m = re.search(r'要：\}[^\n]*\n(?:\s*\\end\{center\}\s*\n)?(.*?)\\par\s*\n', text, re.S)
        if m:
            fm.abstract = [p.strip() for p in re.split(r'\n\s*\n', m.group(1).strip()) if p.strip()]
        m = re.search(r'关键词：\}(.*?)\\par', text)
        if m:
            fm.keywords = [k.strip() for k in re.split(r'\\quad', m.group(1)) if k.strip()]
    return fm


def convert_pdf_figures(engine: str, sections: list[Path], paper: Path, dpi: int = 300) -> list[Path]:
    """把章节里引用的 PDF 图转成 PNG（同目录同名，写入 dpi 元数据以保持物理尺寸）。

    Typst 图片路径相对于所在文件；LaTeX 图片路径相对于 main.tex 所在目录。
    """
    try:
        import pymupdf  # type: ignore
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError:
            pymupdf = None
    pattern = IMAGE_TYP if engine == "typst" else IMAGE_TEX
    made: list[Path] = []
    for sec in sections:
        for m in pattern.finditer(sec.read_text(encoding="utf-8")):
            base = sec.parent if engine == "typst" else paper
            pdf = (base / m.group(1)).resolve()
            png = pdf.with_suffix(".png")
            if not pdf.exists():
                print(f"[warn] 图片不存在: {pdf}")
                continue
            if png.exists() and png.stat().st_mtime >= pdf.stat().st_mtime:
                continue
            if pymupdf is not None:
                doc = pymupdf.open(pdf)
                pix = doc[0].get_pixmap(dpi=dpi)
                pix.set_dpi(dpi, dpi)
                pix.save(png)
                doc.close()
            elif shutil.which("pdftoppm"):
                run(["pdftoppm", "-png", "-r", str(dpi), "-singlefile", str(pdf), str(png.with_suffix(""))])
            else:
                print(f"[warn] 无 pymupdf/pdftoppm，无法转换 {pdf}")
                continue
            made.append(png)
    return made


def resolve_typst_refs(sections: list[Path], paper: Path) -> tuple[Path, list[Path]]:
    """pandoc 不解析 Typst 的 @label 交叉引用，会原样输出 [label]。

    按 Typst 默认规则（图/表/公式各自全局连续编号）预先把 @fig-x / @tbl-x / @eq-x
    替换成“图N”/“表N”/“式(N)”。改写后的章节写到 paper 的同级目录 _docx_shadow/ 下
    （保持目录深度，图片相对路径不变）。返回 (shadow 根目录, 影子章节列表)。
    """
    numbers: dict[str, str] = {}
    counters = {"图": 0, "表": 0}
    eq_n = 0
    for sec in sections:
        text = sec.read_text(encoding="utf-8")
        text = re.sub(r"```.*?```", "", text, flags=re.S)
        events: list[tuple[int, str, str]] = []
        for m in DISPLAY_EQ_TYP.finditer(text):
            events.append((m.start(), "eq", text[m.end():m.end() + 80]))
        for m in LABEL_TYP.finditer(text):
            events.append((m.start(), "label", m.group(1)))
        events.sort()
        pending_eq: int | None = None
        for pos, kind, payload in events:
            if kind == "eq":
                eq_n += 1
                pending_eq = eq_n
                continue
            lab = payload
            if lab.startswith("sec"):
                continue
            if lab.startswith("eq") and pending_eq is not None:
                numbers[lab] = f"式({pending_eq})"
                pending_eq = None
                continue
            if lab.startswith("eq"):
                continue
            # 标签所属图元：向前找最近的 fig( / tbl( / figure( / table( 调用
            head = text[max(0, pos - 4000):pos]
            last = max(head.rfind("tbl("), head.rfind("kind: table"), head.rfind("#table("))
            last_fig = max(head.rfind("fig("), head.rfind("figure("), head.rfind("image("))
            k = "表" if (lab.startswith("tbl") or last > last_fig) else "图"
            counters[k] += 1
            numbers[lab] = f"{k}{counters[k]}"
    shadow = paper.parent / "_docx_shadow"
    if shadow.exists():
        shutil.rmtree(shadow)
    for src in paper.rglob("*"):
        if src.is_dir() or src.name.startswith("_docx") or "docx_render" in src.parts:
            continue
        dst = shadow / src.relative_to(paper)
        dst.parent.mkdir(parents=True, exist_ok=True)
        if src.suffix == ".typ":
            text = src.read_text(encoding="utf-8")
            text = re.sub(r'@([A-Za-z][\w:.-]*)', lambda m: numbers.get(m.group(1), m.group(0)), text)
            text = LABEL_TYP.sub("", text)
            dst.write_text(text, encoding="utf-8")
        elif src.suffix.lower() in (".png", ".jpg", ".jpeg", ".svg", ".gif"):
            # paper/ 内部的图片（含刚转出的 PNG）也要能从影子目录按相对路径找到
            shutil.copy(src, dst)
    return shadow, [shadow / s.relative_to(paper) for s in sections]


def front_matter_source(engine: str, fm: FrontMatter) -> str:
    """用排版源语言写出标题 / 摘要 / 关键词，交给 pandoc 一起转，这样摘要里的公式也能保留。

    三个零宽标记在 style_front_matter() 中被识别并替换成官方摘要页格式。
    """
    title = fm.title or "[论文标题]"
    abstract = fm.abstract or ["[中文摘要内容]"]
    keywords = fm.keywords or ["[关键词]"]
    if engine == "typst":
        return (
            f"{TITLE_MARK}{title}\n\n"
            f"{ABSTRACT_MARK}摘\u3000要：\n\n"
            + "\n\n".join(abstract)
            + f"\n\n{KEYWORDS_MARK}关键词：" + "\u3000\u3000".join(keywords) + "\n\n#pagebreak()\n\n"
        )
    return (
        f"{TITLE_MARK}{title}\n\n"
        f"{ABSTRACT_MARK}摘\u3000要：\n\n"
        + "\n\n".join(abstract)
        + f"\n\n{KEYWORDS_MARK}关键词：" + "\u3000\u3000".join(keywords) + "\n\n\\begin{pagebreak}\n\\end{pagebreak}\n\n"
    )


# --------------------------------------------------------------------------- OOXML 小工具
# Word 对子元素顺序敏感（乱序会报“无法读取的内容”），按 ECMA-376 schema 顺序插入。
_ORDER: dict[str, list[str]] = {
    "w:pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"],
    "w:rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "w:tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
                "tblLook", "tblCaption", "tblDescription"],
    "w:tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"],
    "w:tblBorders": ["top", "start", "left", "bottom", "end", "right", "insideH", "insideV"],
    "w:tcBorders": ["top", "start", "left", "bottom", "end", "right", "insideH", "insideV", "tl2br", "tr2bl"],
    "w:pBdr": ["top", "left", "bottom", "right", "between", "bar"],
    "w:tblCellMar": ["top", "start", "left", "bottom", "end", "right"],
    "w:sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                 "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                 "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid"],
    "w:style": ["name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden", "uiPriority", "semiHidden",
                "unhideWhenUsed", "qFormat", "locked", "personal", "personalCompose", "personalReply", "rsid",
                "pPr", "rPr", "tblPr", "trPr", "tcPr", "tblStylePr"],
    "w:p": ["pPr"],
    "w:tc": ["tcPr"],
    "w:tbl": ["tblPr", "tblGrid", "tr"],
    "w:rPrDefault": ["rPr"],
}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.split(":")[-1]


def _insert_ordered(parent, el) -> None:
    order = _ORDER.get(f"w:{_local(parent.tag)}")
    name = _local(el.tag)
    if order is None or name not in order:
        parent.append(el)
        return
    rank = order.index(name)
    for sib in parent:
        sib_name = _local(sib.tag)
        if sib_name in order and order.index(sib_name) > rank:
            sib.addprevious(el)
            return
    parent.append(el)


def _child(parent, tag: str, **attrs):
    """取或按 schema 顺序新建子元素（属性用 w: 前缀名，如 val="x" -> w:val）。"""
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        _insert_ordered(parent, el)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _set_fonts(rpr, west: str = TNR, east: str = SONG) -> None:
    rf = _child(rpr, "w:rFonts")
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn(f"w:{a}"), west)
    rf.set(qn("w:eastAsia"), east)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{a}"), None)


def _set_size(rpr, half_pts: int) -> None:
    _child(rpr, "w:sz", val=half_pts)
    _child(rpr, "w:szCs", val=half_pts)


def _set_bold(rpr, bold: bool) -> None:
    for tag in ("w:b", "w:bCs"):
        el = rpr.find(qn(tag))
        if bold:
            if el is None:
                el = OxmlElement(tag)
                _insert_ordered(rpr, el)
            el.attrib.pop(qn("w:val"), None)
        elif el is not None:
            rpr.remove(el)


def _clear(rpr_or_ppr, *tags: str) -> None:
    for t in tags:
        for el in rpr_or_ppr.findall(qn(t)):
            rpr_or_ppr.remove(el)


def _ppr(el):
    return _child(el, "w:pPr")


def _para_props(ppr, *, jc: str | None = None, first_line_chars: int | None = None,
                left: int | None = None, hanging: int | None = None,
                before: int | None = None, after: int | None = None,
                line: int | None = None, line_rule: str | None = None,
                keep_next: bool | None = None, keep_lines: bool | None = None,
                outline: int | None = None) -> None:
    if keep_next:
        _child(ppr, "w:keepNext")
    if keep_lines:
        _child(ppr, "w:keepLines")
    if before is not None or after is not None or line is not None:
        sp = _child(ppr, "w:spacing")
        if before is not None:
            sp.set(qn("w:before"), str(before))
            sp.attrib.pop(qn("w:beforeLines"), None)
        if after is not None:
            sp.set(qn("w:after"), str(after))
            sp.attrib.pop(qn("w:afterLines"), None)
        if line is not None:
            sp.set(qn("w:line"), str(line))
            sp.set(qn("w:lineRule"), line_rule or "auto")
    if first_line_chars is not None or left is not None or hanging is not None:
        ind = _child(ppr, "w:ind")
        for a in ("firstLine", "firstLineChars", "left", "leftChars", "hanging", "hangingChars", "start", "startChars"):
            ind.attrib.pop(qn(f"w:{a}"), None)
        if first_line_chars is not None:
            ind.set(qn("w:firstLineChars"), str(first_line_chars))
            ind.set(qn("w:firstLine"), str(first_line_chars * 12 // 5))  # 12 pt 字号下 100 chars = 240 twips
        if left is not None:
            ind.set(qn("w:left"), str(left))
        if hanging is not None:
            ind.set(qn("w:hanging"), str(hanging))
    if jc is not None:
        _child(ppr, "w:jc", val=jc)
    if outline is not None:
        _child(ppr, "w:outlineLvl", val=outline)


def _rpr_of_style(st):
    return _child(st, "w:rPr")


def _ppr_of_style(st):
    return _child(st, "w:pPr")


def _run(text: str, *, west: str = TNR, east: str = SONG, size: int | None = None,
         bold: bool = False, underline: bool = False):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    r.append(rpr)
    _set_fonts(rpr, west, east)
    if bold:
        _set_bold(rpr, True)
    if size:
        _set_size(rpr, size)
    if underline:
        _child(rpr, "w:u", val="single")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tab_run():
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    return r


def _para(style: str | None = None):
    p = OxmlElement("w:p")
    if style:
        _child(_ppr(p), "w:pStyle", val=style)
    return p


def _p_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _p_style(p) -> str:
    ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return ps.get(qn("w:val")) if ps is not None else ""


def _page_break_para():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


# --------------------------------------------------------------------------- 样式（官方正文格式）
def _ensure_style(d, style_id: str, name: str, kind: str = "paragraph", based_on: str | None = None):
    styles_el = d.styles.element
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")).lower() == name.lower():
            return st
    st = OxmlElement("w:style")
    st.set(qn("w:type"), kind)
    st.set(qn("w:styleId"), style_id)
    _child(st, "w:name", val=name)
    if based_on:
        _child(st, "w:basedOn", val=based_on)
    _child(st, "w:qFormat")
    styles_el.append(st)
    return st


def _set_pstyle(p, style_id: str) -> None:
    ppr = _ppr(p)
    _clear(ppr, "w:pStyle")
    _child(ppr, "w:pStyle", val=style_id)


def _find_style(d, *names: str):
    styles_el = d.styles.element
    lowered = {n.lower() for n in names}
    for st in styles_el.findall(qn("w:style")):
        if st.get(qn("w:styleId")) in names:
            return st
        nm = st.find(qn("w:name"))
        if nm is not None and nm.get(qn("w:val")).lower() in lowered:
            return st
    return None


def _style_text(st, *, west: str = TNR, east: str = SONG, size: int | None = None,
                bold: bool | None = None, color_black: bool = True, italic: bool | None = None) -> None:
    rpr = _rpr_of_style(st)
    _set_fonts(rpr, west, east)
    if size is not None:
        _set_size(rpr, size)
    if bold is not None:
        _set_bold(rpr, bold)
    if italic is not None:
        _clear(rpr, "w:i", "w:iCs")
        if italic:
            _child(rpr, "w:i")
    if color_black:
        _clear(rpr, "w:color", "w:shd")
    _child(rpr, "w:kern", val=2)
    _child(rpr, "w:lang", val="en-US", eastAsia="zh-CN")


def apply_body_styles(d) -> None:
    """把输出文档中 pandoc 用到的样式改成官方正文格式（不依赖 reference.docx 的名字对得上）。"""
    # 正文族：Normal / Body Text / First Paragraph（首行缩进 2 字符，两端对齐，单倍行距，段前后 0）
    for names in (("Normal",), ("Body Text", "BodyText"), ("First Paragraph", "FirstParagraph")):
        st = _find_style(d, *names)
        if st is None:
            continue
        _style_text(st, size=24, bold=False, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:spacing", "w:ind", "w:jc")
        _para_props(ppr, jc="both", first_line_chars=200, before=0, after=0, line=240, line_rule="auto")
    # Compact（列表项 / 单元格）：无缩进
    st = _find_style(d, "Compact")
    if st is not None:
        _style_text(st, size=24)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing")
        _para_props(ppr, first_line_chars=0, before=60, after=60)
    # 标题：一级 黑体 14 pt 居中；二级 12 pt 加粗；三级 12 pt
    spec = {
        1: dict(east=HEI, size=28, bold=False, jc="center"),
        2: dict(east=SONG, size=24, bold=True, jc="left"),
        3: dict(east=SONG, size=24, bold=False, jc="left"),
    }
    for lvl, sp in spec.items():
        st = _find_style(d, f"Heading {lvl}", f"Heading{lvl}", f"heading {lvl}")
        if st is None:
            continue
        _style_text(st, east=sp["east"], size=sp["size"], bold=sp["bold"], italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:pBdr", "w:numPr")
        _para_props(ppr, jc=sp["jc"], first_line_chars=0, before=120, after=120, line=240, line_rule="auto",
                    keep_next=True, keep_lines=True, outline=lvl - 1)
    for lvl in range(4, 10):
        st = _find_style(d, f"Heading {lvl}", f"Heading{lvl}")
        if st is not None:
            _style_text(st, size=24, bold=True, italic=False)
    # 图题 / 表题：11 pt 加粗居中，无缩进
    for names, before, after in ((("Image Caption", "ImageCaption"), 60, 120),
                                 (("Table Caption", "TableCaption"), 120, 60),
                                 (("Caption",), 60, 120)):
        st = _find_style(d, *names)
        if st is None:
            continue
        _style_text(st, size=22, bold=True, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc")
        _para_props(ppr, jc="center", first_line_chars=0, before=before, after=after, line=240, line_rule="auto",
                    keep_next=names[0].startswith("Table"), keep_lines=True)
    # 图片段落：居中、无缩进、与图题同页
    for names in (("Captioned Figure", "CaptionedFigure"), ("Figure",)):
        st = _find_style(d, *names)
        if st is None:
            continue
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc")
        _para_props(ppr, jc="center", first_line_chars=0, before=120, after=60, keep_next=True)
    # 代码清单：等宽 10.5 pt，0.5 pt 全框线，单倍行距，无缩进
    st = _find_style(d, "Source Code", "SourceCode")
    if st is not None:
        _style_text(st, west=CODE_FONT, east=SONG, size=21, bold=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:pBdr", "w:shd")
        _para_props(ppr, jc="left", first_line_chars=0, before=0, after=0, line=240, line_rule="auto")
        bdr = _child(ppr, "w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            _child(bdr, f"w:{side}", val="single", sz=4, space=4, color="auto")
    st = _find_style(d, "Verbatim Char", "VerbatimChar")
    if st is not None:
        rpr = _rpr_of_style(st)
        _set_fonts(rpr, CODE_FONT, SONG)
        _set_size(rpr, 21)
    # 参考文献条目：12 pt 两端对齐，固定行距 18 pt，悬挂缩进 482
    st = _ensure_style(d, "References", "References", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="both", left=482, hanging=482, before=0, after=0, line=360, line_rule="exact")
    # 目录
    st = _find_style(d, "TOC Heading", "TOCHeading")
    if st is None:
        st = _ensure_style(d, "TOCHeading", "TOC Heading", based_on="Normal")
    _style_text(st, east=SONG, size=32, bold=True, italic=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:numPr", "w:outlineLvl", "w:pBdr")
    _para_props(ppr, jc="center", first_line_chars=0, before=0, after=120, line=400, line_rule="exact")
    for lvl, left in ((1, 0), (2, 240), (3, 480)):
        st = _find_style(d, f"toc {lvl}", f"TOC{lvl}")
        if st is None:
            st = _ensure_style(d, f"TOC{lvl}", f"toc {lvl}", based_on="Normal")
        _style_text(st, size=21, bold=False, italic=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:tabs")
        _para_props(ppr, jc="both", first_line_chars=0, left=left, before=60, after=60, line=240, line_rule="auto")
        tabs = _child(ppr, "w:tabs")
        _child(tabs, "w:tab", val="right", leader="dot", pos=9060)
    # 页眉 / 页脚：9 pt 居中
    for sid, name in (("Header", "header"), ("Footer", "footer")):
        st = _find_style(d, name, sid)
        if st is None:
            st = _ensure_style(d, sid, name, based_on="Normal")
        _style_text(st, size=18, bold=False)
        ppr = _ppr_of_style(st)
        _clear(ppr, "w:ind", "w:spacing", "w:jc", "w:tabs", "w:pBdr")
        _para_props(ppr, jc="center", first_line_chars=0, before=60, after=60, line=240, line_rule="auto")
    # 摘要页专用样式
    st = _ensure_style(d, "AbstractBody", "Abstract Body", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="both", first_line_chars=200, before=60, after=60, line=240, line_rule="auto")
    # 公式编号 / 表格正文
    st = _ensure_style(d, "TableText", "Table Text", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="center", first_line_chars=0, before=0, after=0, line=360, line_rule="exact")
    st = _ensure_style(d, "Equation", "Equation", based_on="Normal")
    _style_text(st, size=24, bold=False)
    ppr = _ppr_of_style(st)
    _clear(ppr, "w:ind", "w:spacing", "w:jc")
    _para_props(ppr, jc="center", first_line_chars=0, before=60, after=60, line=240, line_rule="auto")
    # 文档默认字体
    dd = d.styles.element.find(qn("w:docDefaults"))
    if dd is not None:
        rpd = dd.find(qn("w:rPrDefault"))
        if rpd is not None:
            rpr = _child(rpd, "w:rPr")
            _set_fonts(rpr, TNR, SONG)
            _set_size(rpr, 24)


def apply_page_setup(d, *, cover_placeholder: bool) -> None:
    """A4、四边 2.5 cm、页眉/页脚距、页脚居中 9 pt 页码；可选空白封面节（页码从摘要页起为 1）。"""
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is None:
        sect = OxmlElement("w:sectPr")
        body.append(sect)
    s = d.sections[-1]
    s.page_width, s.page_height = Twips(PAGE_W), Twips(PAGE_H)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Twips(MARGIN)
    s.header_distance, s.footer_distance = Twips(HEADER_DIST), Twips(FOOTER_DIST)
    s.footer.is_linked_to_previous = False
    s.header.is_linked_to_previous = False
    hp = s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph()
    for r in list(hp._p):
        if r.tag != qn("w:pPr"):
            hp._p.remove(r)
    _set_pstyle(hp._p, "Header")
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    for r in list(fp._p):
        if r.tag != qn("w:pPr"):
            fp._p.remove(r)
    _set_pstyle(fp._p, "Footer")
    _para_props(_ppr(fp._p), jc="center")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE \\* MERGEFORMAT ")
    fld.append(_run("1", size=18))
    fp._p.append(fld)
    _clear(sect, "w:pgNumType", "w:titlePg")
    _child(sect, "w:pgNumType", start=1)

    if cover_placeholder:
        cover = _para()
        _para_props(_ppr(cover), jc="center", first_line_chars=0, before=4000)
        cover.append(_run("[封面：请用当届官方模板封面替换本页，删除本行]", east=HEI, size=28, bold=True))
        brk = _para()
        csect = copy.deepcopy(sect)
        _clear(csect, "w:headerReference", "w:footerReference", "w:pgNumType")
        _child(csect, "w:pgNumType", start=0)
        _ppr(brk).append(csect)
        body.insert(0, brk)
        body.insert(0, cover)


# --------------------------------------------------------------------------- 正文结构后处理
def _iter_body_paras(body):
    return [p for p in body.iterchildren(qn("w:p"))]


def fix_heading_numbers(body) -> list[tuple[int, str]]:
    """pandoc --number-sections 输出 `3<tab>标题` → 一级改为 `三、 标题`，二/三级 `1.1 标题`。

    返回 [(level, 带编号标题文本)]，供目录缓存条目使用。
    """
    headings: list[tuple[int, str]] = []
    for p in body.iter(qn("w:p")):
        st = _p_style(p)
        m = re.fullmatch(r"Heading(\d)", st)
        if not m:
            continue
        lvl = int(m.group(1))
        runs = [r for r in p.findall(qn("w:r"))]
        num_run = next((r for r in runs if r.find(f"{qn('w:rPr')}/{qn('w:rStyle')}") is not None
                        and r.find(f"{qn('w:rPr')}/{qn('w:rStyle')}").get(qn("w:val")) == "SectionNumber"), None)
        if num_run is not None:
            t = num_run.find(qn("w:t"))
            if lvl == 1 and t is not None and t.text and t.text.strip().isdigit():
                t.text = cn_number(int(t.text.strip())) + "、"
            rpr = num_run.find(qn("w:rPr"))
            _set_fonts(rpr, TNR, HEI if lvl == 1 else SONG)
            idx = list(p).index(num_run)
            nxt = p[idx + 1] if idx + 1 < len(p) else None
            if nxt is not None and nxt.tag == qn("w:r") and nxt.find(qn("w:tab")) is not None:
                p.remove(nxt)
                if lvl > 1:
                    t.text = (t.text or "") + " "
        if lvl <= 3:
            headings.append((lvl, _p_text(p).strip()))
    return headings


def style_front_matter(body, fm: FrontMatter) -> None:
    """按官方摘要页格式重排 标题行 / “摘 要：” / 摘要段 / 关键词行，并删除零宽标记。"""
    paras = _iter_body_paras(body)
    title_p = abs_p = kw_p = None
    for p in paras[:12]:
        txt = _p_text(p)
        if TITLE_MARK in txt:
            title_p = p
        elif ABSTRACT_MARK in txt:
            abs_p = p
        elif KEYWORDS_MARK in txt:
            kw_p = p
    if title_p is not None:
        title = _p_text(title_p).replace(TITLE_MARK, "").strip()
        for el in list(title_p):
            if el.tag != qn("w:pPr"):
                title_p.remove(el)
        ppr = _ppr(title_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(title_p, "Normal")
        _para_props(ppr, jc="left", first_line_chars=0, before=120, after=120, line=360, line_rule="auto")
        title_p.append(_run("题\u3000目：", east=LI, size=36))
        title_p.append(_run(title, size=28, underline=True))
    if abs_p is not None:
        for el in list(abs_p):
            if el.tag != qn("w:pPr"):
                abs_p.remove(el)
        ppr = _ppr(abs_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(abs_p, "Normal")
        _para_props(ppr, jc="center", first_line_chars=0, before=120, after=120, line=360, line_rule="auto")
        abs_p.append(_run("摘\u3000要：", east=LI, size=36))
    if abs_p is not None and kw_p is not None:
        start, end = paras.index(abs_p), paras.index(kw_p)
        for p in paras[start + 1:end]:
            ppr = _ppr(p)
            _clear(ppr, "w:ind", "w:spacing")
            _set_pstyle(p, "AbstractBody")
            # “针对问题一，”引导词加粗
            first_r = p.find(qn("w:r"))
            first_t = first_r.find(qn("w:t")) if first_r is not None else None
            if first_t is not None and first_t.text:
                m = LEAD_PHRASE.match(first_t.text)
                if m:
                    first_t.text = first_t.text[m.end():]
                    first_r.addprevious(_run(m.group(1), size=24, bold=True))
    if kw_p is not None:
        keywords = fm.keywords or ["[关键词]"]
        for el in list(kw_p):
            if el.tag != qn("w:pPr"):
                kw_p.remove(el)
        ppr = _ppr(kw_p)
        _clear(ppr, "w:ind", "w:jc", "w:spacing")
        _set_pstyle(kw_p, "Normal")
        _para_props(ppr, jc="left", first_line_chars=0, before=240, after=120, line=360, line_rule="auto")
        kw_p.append(_run("关键词：", east=LI, size=36))
        kw_p.append(_run("\u3000\u3000".join(keywords), size=24))


def insert_toc(body, headings: list[tuple[int, str]], pages: dict[str, int] | None = None) -> None:
    """在第一个一级标题前插入目录页：`目录` 标题 + TOC 域（带静态缓存条目）+ 分页。"""
    first_h1 = next((p for p in body.iter(qn("w:p")) if _p_style(p) == "Heading1"), None)
    if first_h1 is None or not headings:
        return
    anchor = first_h1
    title_p = _para("TOCHeading")
    title_p.append(_run("目录", size=32, bold=True))
    anchor.addprevious(title_p)
    n = len(headings)
    for i, (lvl, text) in enumerate(headings):
        p = _para(f"TOC{lvl}")
        if i == 0:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            fc.set(qn("w:dirty"), "true")
            r.append(fc)
            p.append(r)
            r = OxmlElement("w:r")
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            r.append(it)
            p.append(r)
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "separate")
            r.append(fc)
            p.append(r)
        p.append(_run(text, size=21))
        p.append(_tab_run())
        pg = (pages or {}).get(text)
        p.append(_run(str(pg) if pg else "", size=21))
        if i == n - 1:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            r.append(fc)
            p.append(r)
        anchor.addprevious(p)
    anchor.addprevious(_page_break_para())


EQ_TABLE_MARK = "hwb-equation"


def _is_eq_table(tbl) -> bool:
    cap = tbl.find(f"{qn('w:tblPr')}/{qn('w:tblCaption')}")
    return cap is not None and cap.get(qn("w:val")) == EQ_TABLE_MARK


def _equation_table(omath_para, number: int):
    tbl = OxmlElement("w:tbl")
    tblpr = _child(tbl, "w:tblPr")
    _child(tblpr, "w:tblW", w=5000, type="pct")
    _child(tblpr, "w:jc", val="center")
    bd = _child(tblpr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _child(bd, f"w:{side}", val="nil")
    _child(tblpr, "w:tblLayout", type="fixed")
    mar = _child(tblpr, "w:tblCellMar")
    for side in ("left", "right"):
        _child(mar, f"w:{side}", w=0, type="dxa")
    _child(tblpr, "w:tblLook", val="0000", firstRow=0, lastRow=0, firstColumn=0, lastColumn=0, noHBand=1, noVBand=1)
    _child(tblpr, "w:tblCaption", val=EQ_TABLE_MARK)
    grid = OxmlElement("w:tblGrid")
    tbl.append(grid)
    side_w = TEXT_W // 10
    widths = (side_w, TEXT_W - 2 * side_w, side_w)
    for w in widths:
        _child(grid, "w:gridCol", w=w)
    tr = OxmlElement("w:tr")
    tbl.append(tr)
    for i, w in enumerate(widths):
        tc = OxmlElement("w:tc")
        tcpr = _child(tc, "w:tcPr")
        _child(tcpr, "w:tcW", w=w, type="dxa")
        _child(tcpr, "w:vAlign", val="center")
        p = _para("Equation")
        if i == 1:
            p.append(omath_para)
        elif i == 2:
            _para_props(_ppr(p), jc="right")
            p.append(_run(f"({number})", size=24))
        tc.append(p)
        tr.append(tc)
    return tbl


def number_equations(body) -> int:
    """行间公式（oMathPara）→ 三栏无框表：空 | 公式居中 | (N) 右对齐。返回公式数。"""
    n = 0
    made: list = []
    for omp in list(body.iter(qn("m:oMathPara"))):
        p = omp.getparent()
        if p is None or p.tag != qn("w:p"):
            continue
        # 表格内的公式保持原样
        anc = p.getparent()
        if anc is not None and anc.tag == qn("w:tc"):
            continue
        n += 1
        tbl = _equation_table(omp, n)
        p.addprevious(tbl)
        made.append(tbl)
        if not _p_text(p).strip() and p.find(qn("m:oMathPara")) is None and p.find(qn("w:drawing")) is None:
            p.getparent().remove(p)
    # 相邻两个表在 Word 中会合并，中间垫一个极小空段
    for tbl in made:
        for sib in (tbl.getnext(), tbl.getprevious()):
            if sib is not None and sib.tag == qn("w:tbl"):
                spacer = _para()
                _para_props(_ppr(spacer), before=0, after=0, line=20, line_rule="exact")
                _child(_child(_ppr(spacer), "w:rPr"), "w:sz", val=2)
                tbl.addnext(spacer) if sib is tbl.getnext() else tbl.addprevious(spacer)
    return n


def style_tables(body) -> int:
    """普通表 → 三线表：居中，顶线/底线 1.5 pt，表头下线 0.5 pt，表头加粗，单元格 12 pt 固定行距 18 pt。"""
    n = 0
    for tbl in body.iter(qn("w:tbl")):
        tblpr = tbl.find(qn("w:tblPr"))
        if tblpr is None or _is_eq_table(tbl):
            continue
        if tbl.getparent().tag == qn("w:tc"):
            continue  # 嵌套表
        n += 1
        _child(tblpr, "w:jc", val="center")
        _clear(tblpr, "w:tblBorders")
        bd = _child(tblpr, "w:tblBorders")
        for side, val, sz in (("top", "single", 12), ("left", "nil", 0), ("bottom", "single", 12),
                              ("right", "nil", 0), ("insideH", "nil", 0), ("insideV", "nil", 0)):
            el = _child(bd, f"w:{side}", val=val)
            if val != "nil":
                el.set(qn("w:sz"), str(sz))
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "auto")
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            continue
        header_rows = [r for r in rows if r.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None] or rows[:1]
        last_header = header_rows[-1]
        for tr in rows:
            is_header = tr in header_rows
            for tc in tr.findall(qn("w:tc")):
                tcpr = _child(tc, "w:tcPr")
                if tr is last_header:
                    tb = _child(tcpr, "w:tcBorders")
                    _child(tb, "w:bottom", val="single", sz=4, space=0, color="auto")
                _child(tcpr, "w:vAlign", val="center")
                for p in tc.findall(qn("w:p")):
                    ppr = _ppr(p)
                    jc = ppr.find(qn("w:jc"))
                    jc_val = jc.get(qn("w:val")) if jc is not None else "center"
                    _clear(ppr, "w:ind", "w:spacing", "w:jc")
                    _set_pstyle(p, "TableText")
                    _para_props(ppr, jc=jc_val, first_line_chars=0, before=0, after=0, line=360, line_rule="exact")
                    if is_header:
                        for r in p.findall(qn("w:r")):
                            _set_bold(_child(r, "w:rPr"), True)
    return n


def style_references(body) -> int:
    """`参考文献` 一级标题到下一个一级标题之间的段落套 References 样式。返回条目数。"""
    n = 0
    in_refs = False
    for p in _iter_body_paras(body):
        st = _p_style(p)
        if st == "Heading1":
            in_refs = _p_text(p).strip().startswith("参考文献")
            continue
        if in_refs and st in ("", "Normal", "BodyText", "FirstParagraph", "Compact") and _p_text(p).strip():
            _clear(_ppr(p), "w:ind", "w:spacing")
            _set_pstyle(p, "References")
            n += 1
    return n


def apply_huaweibei_body(docx_path: Path, fm: FrontMatter, *, front_matter: bool, cover_placeholder: bool,
                         page_setup: bool, toc_pages: dict[str, int] | None = None) -> dict:
    d = docx.Document(str(docx_path))
    body = d.element.body
    apply_body_styles(d)
    headings = fix_heading_numbers(body)
    if front_matter:
        style_front_matter(body, fm)
    n_eq = number_equations(body)
    n_tbl = style_tables(body)
    n_ref = style_references(body)
    insert_toc(body, headings, toc_pages)
    if page_setup:
        apply_page_setup(d, cover_placeholder=cover_placeholder)
    d.save(str(docx_path))
    return {"numbered_equations": n_eq, "three_line_tables": n_tbl, "reference_entries": n_ref,
            "toc_entries": [t for _, t in headings]}


# --------------------------------------------------------------------------- reference.docx
def build_reference_docx(dst: Path) -> None:
    """pandoc 默认 reference.docx + A4 页面（pandoc 据此缩放图片）；样式细节在输出文档上统一施加。"""
    raw = run(["pandoc", "-o", str(dst), "--print-default-data-file", "reference.docx"])
    if raw.returncode != 0:
        sys.exit(f"pandoc 生成 reference.docx 失败: {raw.stderr}")
    d = docx.Document(str(dst))
    apply_body_styles(d)
    s = d.sections[0]
    s.page_width, s.page_height = Twips(PAGE_W), Twips(PAGE_H)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Twips(MARGIN)
    s.header_distance, s.footer_distance = Twips(HEADER_DIST), Twips(FOOTER_DIST)
    d.save(str(dst))


# --------------------------------------------------------------------------- 检查
def inspect(docx_path: Path) -> dict:
    z = zipfile.ZipFile(docx_path)
    xml = z.read("word/document.xml").decode("utf-8")
    text = re.sub(r"<[^>]+>", "", xml)
    footers = "".join(z.read(n).decode("utf-8") for n in z.namelist() if re.match(r"word/footer\d*\.xml", n))
    sect = re.findall(r"<w:sectPr.*?</w:sectPr>", xml, flags=re.S)
    last_sect = sect[-1] if sect else ""
    pgsz = re.search(r'<w:pgSz[^>]*w:w="(\d+)"[^>]*w:h="(\d+)"', last_sect) or re.search(r'<w:pgSz[^>]*w:h="(\d+)"[^>]*w:w="(\d+)"', last_sect)
    pgmar = re.search(r"<w:pgMar([^>]*)/>", last_sect)
    margins = dict(re.findall(r'w:(top|bottom|left|right)="(\d+)"', pgmar.group(1))) if pgmar else {}
    styles_used = re.findall(r'<w:pStyle w:val="([^"]+)"', xml)
    fig_paras = styles_used.count("CaptionedFigure") + styles_used.count("Figure")
    n_eq_tables = xml.count(f'<w:tblCaption w:val="{EQ_TABLE_MARK}"')
    return {
        "paragraphs": xml.count("<w:p>") + xml.count("<w:p "),
        "images": sum(1 for n in z.namelist() if n.startswith("word/media/")),
        "equations": xml.count("<m:oMath>") + xml.count("<m:oMath "),
        "display_equations": xml.count("<m:oMathPara>") + xml.count("<m:oMathPara "),
        "tables": xml.count("<w:tbl>") - n_eq_tables,
        "headings": {f"h{l}": styles_used.count(f"Heading{l}") for l in (1, 2, 3)},
        "image_captions": styles_used.count("ImageCaption"),
        "table_captions": styles_used.count("TableCaption"),
        "uncaptioned_figures": styles_used.count("Figure"),
        "figure_paragraphs": fig_paras,
        "page_size_twips": [int(pgsz.group(1)), int(pgsz.group(2))] if pgsz else None,
        "margins_twips": {k: int(v) for k, v in margins.items()},
        "footer_page_field": "PAGE" in footers,
        "toc_field": "TOC \\o" in xml,
        "placeholders": sorted(set(PLACEHOLDER.findall(text))),
        "leaked_internal_names": sorted({w for w in INTERNAL_NAMES if w in text}),
    }


def audit(report: dict, *, expect_page_setup: bool) -> tuple[list[str], list[str]]:
    """返回 (errors, warnings)：errors 在 --strict 下非零退出，warnings 仅提示。"""
    errors: list[str] = []
    warnings: list[str] = []
    if report["placeholders"]:
        errors.append("DOCX 中仍有模板占位符，提交前必须替换: " + "、".join(report["placeholders"]))
    if report["leaked_internal_names"]:
        errors.append("DOCX 中出现工作流内部文件名，提交前必须删除: " + "、".join(report["leaked_internal_names"]))
    if expect_page_setup:
        if sorted(report["page_size_twips"] or []) != sorted([PAGE_W, PAGE_H]):
            errors.append(f"页面尺寸不是 A4: {report['page_size_twips']}")
        if len(report["margins_twips"]) < 4 or any(v != MARGIN for v in report["margins_twips"].values()):
            errors.append(f"页边距不是四边 2.5 cm: {report['margins_twips']}")
        if not report["footer_page_field"]:
            errors.append("页脚缺少 PAGE 页码域")
    if report["headings"]["h1"] == 0:
        errors.append("没有一级标题（正文未按章节组织？）")
    if report["uncaptioned_figures"]:
        warnings.append(f"{report['uncaptioned_figures']} 张图没有图题")
    if report["images"] and report["image_captions"] < report["images"] - report["uncaptioned_figures"]:
        warnings.append("图题数少于图片数（多图并列时可忽略）")
    if report["tables"] and report["table_captions"] < report["tables"]:
        warnings.append(f"表题数({report['table_captions']})少于表格数({report['tables']})")
    if not report["toc_field"]:
        warnings.append("未生成目录")
    return errors, warnings


def docx_to_pdf(docx_path: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("[info] 未找到 soffice，跳过 DOCX->PDF 渲染抽检")
        return None
    # 输出到独立子目录，避免覆盖排版引擎生成的 main.pdf
    outdir = docx_path.parent / "docx_render"
    outdir.mkdir(exist_ok=True)
    r = run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)], timeout=300)
    pdf = outdir / docx_path.with_suffix(".pdf").name
    if r.returncode != 0 or not pdf.exists():
        print(f"[warn] soffice 转换失败: {r.stderr[-500:]}")
        return None
    return pdf


def locate_heading_pages(pdf: Path, headings: list[str], page_offset: int) -> dict[str, int]:
    """在渲染 PDF 里按行精确匹配标题文本，返回 {标题: 显示页码}（跳过目录页的点线条目）。"""
    try:
        import pymupdf  # type: ignore
    except ImportError:
        return {}
    doc = pymupdf.open(pdf)
    norm = lambda s: re.sub(r"\s+", "", s)
    targets = [norm(h) for h in headings]
    found: dict[str, int] = {}
    start_page = 0
    for h, key in zip(headings, targets):
        for pno in range(start_page, doc.page_count):
            lines = [norm(l) for l in doc[pno].get_text().splitlines()]
            if key in lines:
                found[h] = pno + 1 - page_offset
                start_page = pno
                break
    doc.close()
    return found


# --------------------------------------------------------------------------- 主流程
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--paper", default="paper", help="论文源码目录（含 main.typ 或 main.tex）")
    ap.add_argument("--out", default=None, help="输出 DOCX 路径，默认 <paper>/main.docx")
    ap.add_argument("--reference", default=None, help="Word 参考样式文档（当届官方 DOCX 模板可直接传入；页面设置以其为准）")
    ap.add_argument("--title", default=None)
    ap.add_argument("--abstract-file", default=None, help="摘要纯文本文件，空行分段")
    ap.add_argument("--keywords", default=None, help="分号分隔")
    ap.add_argument("--no-front-matter", action="store_true", help="不生成题目/摘要/关键词页（例如官方模板已含）")
    ap.add_argument("--cover-placeholder", action="store_true", help="最前面加一页空白封面节（页码 0，不显示），摘要页从 1 起")
    ap.add_argument("--figure-dpi", type=int, default=300, help="PDF 图转 PNG 的分辨率（默认 300）")
    ap.add_argument("--pdf", action="store_true", help="用 soffice 顺带导出 PDF 做渲染抽检，并回填目录页码")
    ap.add_argument("--keep-entry", action="store_true", help="pandoc 失败时保留 _docx_body.* 中间文件")
    ap.add_argument("--strict", action="store_true", help="发现占位符、内部文件名泄露或格式审计问题时以非零退出（终稿检查用）")
    args = ap.parse_args()

    if not shutil.which("pandoc"):
        sys.exit("缺少 pandoc，请先运行 scripts/setup_env.sh 或 @skills:doctor")

    paper = Path(args.paper).resolve()
    engine, main_file = detect_engine(paper)
    out = Path(args.out).resolve() if args.out else paper / "main.docx"
    sections = collect_sections(engine, main_file)
    if not sections:
        sys.exit("没有找到任何章节文件（#include / \\input）")
    print(f"[info] 引擎={engine}，章节 {len(sections)} 个")

    fm = extract_front_matter(engine, main_file)
    if args.title:
        fm.title = args.title
    if args.abstract_file:
        fm.abstract = [p.strip() for p in re.split(r"\n\s*\n", Path(args.abstract_file).read_text(encoding="utf-8")) if p.strip()]
    if args.keywords:
        fm.keywords = [k.strip() for k in re.split(r"[;；]", args.keywords) if k.strip()]
    print(f"[info] 标题={fm.title!r} 摘要段数={len(fm.abstract)} 关键词={fm.keywords}")

    pngs = convert_pdf_figures(engine, sections, paper, dpi=args.figure_dpi)
    if pngs:
        print(f"[info] 生成 PNG 图 {len(pngs)} 张")

    with tempfile.TemporaryDirectory() as td:
        tdp = Path(td)
        lua = tdp / "hwb_docx.lua"
        lua.write_text(PDF2PNG_LUA, encoding="utf-8")
        ref = Path(args.reference).resolve() if args.reference else tdp / "reference.docx"
        if not args.reference:
            build_reference_docx(ref)

        # 正文入口：摘要页 + 章节，不带封面 / 手写目录等 Word 不需要的排版代码
        head = "" if args.no_front_matter else front_matter_source(engine, fm)
        shadow: Path | None = None
        work = paper
        if engine == "typst":
            shadow, sections = resolve_typst_refs(sections, paper)
            work = shadow
            entry = work / "_docx_body.typ"
            entry.write_text(head + "".join(f'#include("{s.relative_to(work).as_posix()}")\n' for s in sections), encoding="utf-8")
            fmt = "typst"
        else:
            entry = paper / "_docx_body.tex"
            entry.write_text(
                "\\documentclass{ctexart}\n\\usepackage{graphicx,amsmath,amssymb,booktabs,float}\n"
                "\\newcommand{\\threelinetable}[4]{\\begin{table}[H]\\centering\\caption{#1}"
                "\\begin{tabular}{#2}\\toprule #3 \\\\ \\midrule #4 \\\\ \\bottomrule\\end{tabular}\\end{table}}\n"
                "\\begin{document}\n"
                + head
                + "".join(f"\\input{{{s.relative_to(paper).as_posix()}}}\n" for s in sections)
                + "\\end{document}\n",
                encoding="utf-8",
            )
            fmt = "latex"
        try:
            cmd = [
                "pandoc", str(entry), "-f", fmt, "-t", "docx",
                "--number-sections",
                "--reference-doc", str(ref),
                "--lua-filter", str(lua),
                "--resource-path", str(work),
                "-o", str(out),
            ]
            r = run(cmd, cwd=str(work))
        finally:
            if r.returncode == 0 or not args.keep_entry:
                entry.unlink(missing_ok=True)
                if shadow is not None:
                    shutil.rmtree(shadow, ignore_errors=True)
        if r.returncode != 0:
            hint = f"（已保留 {entry} 供排查）" if args.keep_entry else "（加 --keep-entry 可保留中间文件排查）"
            sys.exit(f"pandoc 失败{hint}:\n{r.stderr}")
        if r.stderr.strip():
            print("[pandoc]", r.stderr.strip()[-800:])

    raw_docx = out.with_name(out.stem + "_pandoc_raw.docx")
    shutil.copy(out, raw_docx)
    page_setup = not args.reference
    body_report = apply_huaweibei_body(out, fm, front_matter=not args.no_front_matter,
                                       cover_placeholder=args.cover_placeholder, page_setup=page_setup)

    report = inspect(out)
    report.update({k: v for k, v in body_report.items() if k != "toc_entries"})
    report["docx"] = str(out)
    if args.pdf:
        pdf = docx_to_pdf(out)
        if pdf:
            pages = locate_heading_pages(pdf, body_report["toc_entries"], 1 if args.cover_placeholder else 0)
            if pages:
                # 回填目录页码后重新渲染
                shutil.copy(raw_docx, out)
                apply_huaweibei_body(out, fm, front_matter=not args.no_front_matter,
                                     cover_placeholder=args.cover_placeholder, page_setup=page_setup, toc_pages=pages)
                pdf = docx_to_pdf(out) or pdf
                report["toc_pages_filled"] = len(pages)
            report["pdf"] = str(pdf)
            try:
                import pymupdf  # type: ignore

                report["pdf_pages"] = pymupdf.open(pdf).page_count
            except Exception:
                pass
    raw_docx.unlink(missing_ok=True)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    errors, warnings = audit(report, expect_page_setup=page_setup)
    for w in warnings:
        print(f"[warn] {w}")
    for e in errors:
        print(f"[error] {e}")
    if errors and args.strict:
        sys.exit(2)


if __name__ == "__main__":
    main()
